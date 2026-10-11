"""results.json -> an Excel workbook, coloured by who decided each cell (MINH_PLAN §B1).

Sheets: Summary (if any), Results, one per aggregate, Stats, Dropped (if any). The
Results sheet keeps `__by` / `__confidence` / `__low_confidence` as hidden columns so a
reviewer can unhide them and filter, while the colours tell the story at a glance.
"""

from __future__ import annotations

import io
import logging
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from foreman.contracts import ResultsFile
from foreman.format.charts import bars_of, find_charts, title_of
from foreman.format.colors import FILL, HEADER_FILL, LABEL, LOW_CONFIDENCE_TEXT, MARK
from foreman.format.table import (
    cell_text,
    columns,
    decided_by,
    is_low,
    ordered_items,
    output_fields,
)

log = logging.getLogger(__name__)

BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor=HEADER_FILL)
FILLS = {k: PatternFill("solid", fgColor=v) for k, v in FILL.items()}
LOW_FONT = Font(color=LOW_CONFIDENCE_TEXT, bold=True)
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
MAX_TEXT_WIDTH = 80


def _sheet_name(wb: Workbook, key: str) -> str:
    base = _BAD_SHEET_CHARS.sub(" ", title_of(key))[:28] or "Sheet"
    name, n = base, 2
    while name in wb.sheetnames:
        name, n = f"{base[:25]} {n}", n + 1
    return name


def _header(ws: Worksheet, row: int, names: list[str]) -> None:
    for c, name in enumerate(names, start=1):
        cell = ws.cell(row=row, column=c, value=name)
        cell.font, cell.fill = BOLD, HEAD_FILL


def _autosize(ws: Worksheet) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                n = len(str(cell.value))
                widths[cell.column] = max(widths.get(cell.column, 0), min(n, MAX_TEXT_WIDTH))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(8, w + 2)


def _results_sheet(wb: Workbook, results: ResultsFile) -> None:
    ws = wb.create_sheet("Results")
    cols = columns(results)
    decided = [c for c in cols if c in output_fields(results)]
    hidden = [f"{c}__{s}" for c in decided for s in ("by", "confidence", "low_confidence")]

    ws.cell(row=1, column=1, value="Decided by:").font = BOLD
    for n, key in enumerate(("jev", "llm", "code"), start=2):
        cell = ws.cell(row=1, column=n, value=LABEL[key])
        cell.fill = FILLS[key]
    low = ws.cell(row=1, column=5, value="Low confidence")
    low.font = LOW_FONT

    names = ["id", "text", *cols, *hidden]
    _header(ws, 2, names)
    for r, item in enumerate(ordered_items(results), start=3):
        ws.cell(row=r, column=1, value=item.id)
        ws.cell(row=r, column=2, value=item.text).alignment = Alignment(wrap_text=False)
        for c, field in enumerate(cols, start=3):
            cell = ws.cell(row=r, column=c, value=cell_text(item.fields.get(field)))
            by = decided_by(item, field)
            if by:
                cell.fill = FILLS[by]
            if is_low(item, field):
                cell.font = LOW_FONT
        for c, field in enumerate(hidden, start=3 + len(cols)):
            ws.cell(row=r, column=c, value=cell_text(item.fields.get(field)))
    _autosize(ws)
    ws.column_dimensions["B"].width = 60
    for c in range(3 + len(cols), 3 + len(cols) + len(hidden)):
        ws.column_dimensions[get_column_letter(c)].hidden = True
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(names))}{max(2, ws.max_row)}"


def _bar_chart(ws: Worksheet, title: str, n: int, anchor: str, color: str | None = None) -> None:
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = title
    chart.legend = None
    chart.height = max(6, 0.6 * n + 2)
    chart.width = 16
    data = Reference(ws, min_col=2, min_row=1, max_row=n + 1)
    cats = Reference(ws, min_col=1, min_row=2, max_row=n + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.majorGridlines = None
    chart.x_axis.scaling.orientation = "maxMin"  # largest bar on top, as in the table
    chart.series[0].graphicalProperties.solidFill = color or MARK["jev"]
    chart.series[0].graphicalProperties.line.noFill = True
    ws.add_chart(chart, anchor)


def _aggregate_sheet(wb: Workbook, key: str, value: Any, chart: bool) -> None:
    ws = wb.create_sheet(_sheet_name(wb, key))
    bars = bars_of(value)
    if isinstance(value, dict):
        _header(ws, 1, [title_of(key), "value"])
        rows = sorted(value.items(), key=lambda kv: -kv[1]) if bars else value.items()
        for r, (k, v) in enumerate(rows, start=2):
            ws.cell(row=r, column=1, value=str(k))
            ws.cell(row=r, column=2, value=cell_text(v))
    elif isinstance(value, list) and value and all(isinstance(r, dict) for r in value):
        keys: dict[str, None] = {}
        for row in value:
            keys.update(dict.fromkeys(row))
        names = list(keys)
        if bars:  # label + count first, so the chart reads columns A and B
            label = next(k for k in names if isinstance(value[0].get(k), str))
            names = [label, "count", *[k for k in names if k not in (label, "count")]]
            value = sorted(value, key=lambda r: -r["count"])
        _header(ws, 1, names)
        for r, row in enumerate(value, start=2):
            for c, k in enumerate(names, start=1):
                ws.cell(row=r, column=c, value=cell_text(row.get(k)))
    elif isinstance(value, list):
        _header(ws, 1, [title_of(key)])
        for r, v in enumerate(value, start=2):
            ws.cell(row=r, column=1, value=cell_text(v))
    else:
        _header(ws, 1, [title_of(key)])
        ws.cell(row=2, column=1, value=cell_text(value))
    _autosize(ws)
    if bars and chart:
        _bar_chart(ws, title_of(key), len(bars), "E2")


def _stats_sheet(wb: Workbook, results: ResultsFile) -> None:
    ws = wb.create_sheet("Stats")
    st = results.stats
    _header(ws, 1, ["decided by", "decisions"])
    for r, key in enumerate(("jev", "llm", "code"), start=2):
        ws.cell(row=r, column=1, value=LABEL[key]).fill = FILLS[key]
        ws.cell(row=r, column=2, value=st.by.get(key, 0))
    _bar_chart(ws, "Decisions by Jev / LLM / code", 3, "E2")
    rows = [
        ("decisions_total", st.decisions_total),
        ("low_confidence", st.low_confidence),
        ("jev_cost_usd", st.jev_cost_usd),
        ("llm_cost_usd", st.llm_cost_usd),
        ("total_cost_usd", st.jev_cost_usd + st.llm_cost_usd),
        ("wall_time_s", st.wall_time_s),
        ("jev_calls", st.jev_calls),
        ("llm_calls", st.llm_calls),
        ("threshold", results.threshold),
        ("run_id", results.run_id),
        ("job", results.job),
        ("started_at", results.started_at),
        ("finished_at", results.finished_at),
    ]
    start = 6
    _header(ws, start, ["stat", "value"])
    for r, (k, v) in enumerate(rows, start=start + 1):
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=v)
    if results.node_reports:
        top = start + len(rows) + 2
        names = ["node", "tool", "items_in", "items_out", "jev_decisions", "llm_decisions",
                 "code_ops", "low_confidence", "duration_ms", "notes"]  # fmt: skip
        _header(ws, top, names)
        for r, rep in enumerate(results.node_reports, start=top + 1):
            data = rep.model_dump()
            for c, k in enumerate(names, start=1):
                v = "; ".join(data[k]) if k == "notes" else data[k]
                ws.cell(row=r, column=c, value=v)
    _autosize(ws)


def _dropped_sheet(wb: Workbook, results: ResultsFile) -> None:
    ws = wb.create_sheet("Dropped")
    _header(ws, 1, ["id", "reason", "text"])
    for r, d in enumerate(results.dropped, start=2):
        ws.cell(row=r, column=1, value=d.id)
        ws.cell(row=r, column=2, value=d.reason)
        ws.cell(row=r, column=3, value=d.text)
    _autosize(ws)


def _summary_sheet(wb: Workbook, results: ResultsFile, summary: str) -> None:
    ws = wb.create_sheet("Summary")
    ws["A1"] = "Summary"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = summary
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 100
    ws.row_dimensions[2].height = max(60, 15 * (len(summary) // 95 + 1))
    ws["A4"] = f"Job: {results.job}"
    ws["A5"] = f"Run {results.run_id}, threshold {results.threshold}"


def to_xlsx(results: ResultsFile, summary: str | None = None) -> bytes:
    """The workbook as bytes."""
    wb = Workbook()
    wb.remove(wb.active)
    if summary:
        _summary_sheet(wb, results, summary)
    _results_sheet(wb, results)
    charted = {c.key for c in find_charts(results)}
    for key, value in results.aggregates.items():
        if key == "summary":
            continue
        # theme_groups and theme_counts hold the same bars: one chart, two sheets.
        _aggregate_sheet(wb, key, value, chart=key in charted)
    _stats_sheet(wb, results)
    if results.dropped:
        _dropped_sheet(wb, results)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
