"""results.json -> a Word report (MINH_PLAN §B3): summary, charts, one table per aggregate.

Word is what a manager forwards; it carries the same charts and colours as the pdf.
"""

from __future__ import annotations

import io
import logging
from typing import Any

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from foreman.contracts import ResultsFile
from foreman.format.charts import bars_of, find_charts, render_png, title_of
from foreman.format.colors import MUTED, hex_rgb
from foreman.format.table import cell_text

log = logging.getLogger(__name__)

MAX_ROWS = 30


def _shade(cell: Any, hex_fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    props.append(shd)


def _table(document: Any, header: list[str], rows: list[list[Any]]) -> None:
    t = document.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for cell, name in zip(t.rows[0].cells, header, strict=True):
        cell.text = str(name)
    for row in rows[:MAX_ROWS]:
        cells = t.add_row().cells
        for cell, v in zip(cells, row, strict=True):
            cell.text = "" if v is None else str(cell_text(v))
    if len(rows) > MAX_ROWS:
        p = document.add_paragraph(f"… and {len(rows) - MAX_ROWS} more rows (see the xlsx export).")
        p.runs[0].font.color.rgb = RGBColor(*hex_rgb(MUTED))


def _aggregate_rows(value: Any) -> tuple[list[str], list[list[Any]]] | None:
    if isinstance(value, dict):
        items = value.items()
        if bars_of(value):
            items = sorted(items, key=lambda kv: -kv[1])
        return ["key", "value"], [[k, v] for k, v in items]
    if isinstance(value, list) and value and all(isinstance(r, dict) for r in value):
        keys: dict[str, None] = {}
        for r in value:
            keys.update(dict.fromkeys(k for k in r if k not in {"example_ids", "member_ids"}))
        return list(keys), [[r.get(k) for k in keys] for r in value]
    return None


def to_docx(results: ResultsFile, summary: str | None = None) -> bytes:
    """The report as docx bytes."""
    d = docx.Document()
    d.styles["Normal"].font.size = Pt(10.5)
    d.add_heading("Foreman run report", level=0)
    st = results.stats
    meta = d.add_paragraph(
        f"Job: {results.job}\nRun {results.run_id} · threshold {results.threshold} · "
        f"{st.decisions_total:,} decisions (Jev {st.by.get('jev', 0):,}, LLM {st.by.get('llm', 0):,}, "
        f"code {st.by.get('code', 0):,}) · ${st.jev_cost_usd + st.llm_cost_usd:.4f} · "
        f"{st.wall_time_s:.1f} s"
    )
    for run in meta.runs:
        run.font.color.rgb = RGBColor(*hex_rgb(MUTED))
    if summary:
        d.add_heading("Summary", level=1)
        d.add_paragraph(summary)

    d.add_heading("Charts", level=1)
    for chart in find_charts(results):
        d.add_picture(io.BytesIO(render_png(chart, 6.0)), width=Inches(6.0))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT

    d.add_heading("Aggregates", level=1)
    for key, value in results.aggregates.items():
        if key in {"summary", "escalation_curve", "rank_order"}:
            continue
        shaped = _aggregate_rows(value)
        if shaped is None:
            continue
        d.add_heading(title_of(key), level=2)
        _table(d, *shaped)

    if results.node_reports:
        d.add_heading("Steps", level=1)
        _table(
            d,
            ["node", "tool", "items in", "items out", "Jev", "LLM", "code", "seconds"],
            [
                [r.node, r.tool, r.items_in, r.items_out, r.jev_decisions, r.llm_decisions,
                 r.code_ops, round(r.duration_ms / 1000, 1)]
                for r in results.node_reports
            ],
        )  # fmt: skip
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
