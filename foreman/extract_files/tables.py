"""csv and xlsx -> one item per row (MINH_PLAN §A1).

Ids are `row-<n>` with the spreadsheet row number, so the first data row is `row-2`; a
user can find any result back in their own sheet. `fields` keeps every column, typed, so
`calculate` can do math on numbers in code.
"""

from __future__ import annotations

import csv
import io
import logging
import math
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from foreman.contracts import ItemsFile
from foreman.core.state import ItemState
from foreman.extract_files.detect import ExtractError, decode_text

log = logging.getLogger(__name__)

MIN_TEXT_AVG = 20  # a "free text" column averages at least this many characters
_INT = re.compile(r"-?\d{1,15}")
_LEADING_ZERO = re.compile(r"-?0\d")
_FLOAT = re.compile(r"-?(\d+\.\d*|\.\d+)([eE][-+]?\d+)?")


def clean_value(v: Any) -> Any:
    """A cell -> a JSON-safe value: numbers stay numbers, dates become ISO, blanks None."""
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return None
        return int(v) if v.is_integer() and abs(v) < 1e15 else v
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0) else v.isoformat(timespec="seconds")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.isoformat()
    s = str(v).strip()
    if not s:
        return None
    return _typed_text(s)


def _typed_text(s: str) -> Any:
    """csv cells arrive as text; numbers in them become numbers.

    Codes with a leading zero ("02134") stay text, since they are ids, not quantities.
    """
    if _INT.fullmatch(s) and not _LEADING_ZERO.match(s):
        return int(s)
    if _FLOAT.fullmatch(s):
        return float(s)
    return s


def text_column(header: list[str], rows: list[dict[str, Any]]) -> str | None:
    """The string column with the longest average length, if it averages >= 20 chars."""
    best, best_len = None, 0.0
    for col in header:
        values = [r[col] for r in rows if isinstance(r.get(col), str)]
        if not values:
            continue
        avg = sum(len(v) for v in values) / len(values)
        if avg > best_len:
            best, best_len = col, avg
    return best if best_len >= MIN_TEXT_AVG else None


def rows_to_items(
    header: list[str],
    rows: list[dict[str, Any]],
    *,
    source: dict[str, Any],
    sheet: str | None = None,
) -> ItemsFile:
    """Typed rows (row 1 is the header) -> ItemsFile. Empty rows are skipped."""
    col = text_column(header, rows)
    items: list[ItemState] = []
    for n, row in enumerate(rows, start=2):
        fields = {k: v for k in header if (v := row.get(k)) is not None}
        if not fields:
            continue
        text = row.get(col) if col else None
        if not isinstance(text, str) or not text.strip():
            text = "; ".join(f"{k}: {v}" for k, v in fields.items())
        meta: dict[str, Any] = {"row": n}
        if sheet is not None:
            meta["sheet"] = sheet
        items.append(ItemState(id=f"row-{n}", text=text.strip(), fields=fields, meta=meta))
    return ItemsFile(source={**source, "rows": len(rows), "text_column": col}, items=items)


def _unique_header(raw: list[Any]) -> list[str]:
    """Blank headers become `column_<n>`, repeats get a suffix, so no column is lost."""
    out: list[str] = []
    for i, h in enumerate(raw, start=1):
        name = str(h).strip() if h is not None and str(h).strip() else f"column_{i}"
        base, k = name, 2
        while name in out:
            name, k = f"{base}_{k}", k + 1
        out.append(name)
    return out


def read_csv(path: Path) -> ItemsFile:
    """csv (or tsv / semicolon) with a header row."""
    text = decode_text(path.read_bytes(), path.name)
    sample = "\n".join(text.splitlines()[:20])
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    table = [r for r in reader]
    if not table or not any(c.strip() for c in table[0]):
        raise ExtractError(f"{path.name} has no header row.")
    header = _unique_header(table[0])
    rows = [
        {h: clean_value(r[i]) if i < len(r) else None for i, h in enumerate(header)}
        for r in table[1:]
    ]
    return rows_to_items(header, rows, source={"filename": path.name, "kind": "csv", "sheet": None})


def read_xlsx(path: Path, sheet: str | None = None) -> ItemsFile:
    """The first sheet (or `sheet`) of an xlsx; the header is the sheet's first row."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # corrupt zip, wrong content
        raise ExtractError(f"{path.name} is not a valid xlsx file: {e}") from e
    try:
        if sheet is not None and sheet not in wb.sheetnames:
            raise ExtractError(
                f"{path.name} has no sheet {sheet!r}; sheets: {', '.join(wb.sheetnames)}"
            )
        ws = wb[sheet] if sheet is not None else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        first = next(it, None)
        if first is None or all(c is None or not str(c).strip() for c in first):
            raise ExtractError(f"{path.name} (sheet {ws.title}) has no header row.")
        header = _unique_header(list(first))
        rows = [
            {h: clean_value(r[i]) if i < len(r) else None for i, h in enumerate(header)} for r in it
        ]
        title = ws.title
    finally:
        wb.close()
    while rows and all(v is None for v in rows[-1].values()):
        rows.pop()  # xlsx files often carry formatted but empty rows at the end
    source = {"filename": path.name, "kind": "xlsx", "sheet": title}
    return rows_to_items(header, rows, source=source, sheet=title)
