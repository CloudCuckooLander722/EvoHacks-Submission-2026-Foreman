"""summarize() and export(): numbers checked in code, sheets and colours as specified."""

from io import BytesIO

import pytest

from foreman.contracts import DroppedItem, NodeReport, ResultsFile
from foreman.core.context import make_context
from foreman.core.ledger import RunStats
from foreman.core.state import ItemState
from foreman.format import export, summarize, template_summary
from foreman.format.charts import find_charts
from foreman.format.colors import FILL
from foreman.format.summary import invented_numbers, summary_input
from foreman.testing.fakes import FakeLLM


def _item(n: int, type_: str, by: str, theme: str, low: bool = False) -> ItemState:
    fields = {
        "rating": n % 5 + 1,
        "type": type_,
        "type__confidence": 0.9 if by == "jev" else 0.6,
        "type__by": by,
        "type__probs": {type_: 0.9},
        "theme": theme,
        "theme__by": "jev",
        "theme__confidence": 0.95,
        "rank": n,
        "rank__by": "code",
    }
    if low:
        fields["type__low_confidence"] = True
    return ItemState(id=f"row-{n}", text=f"review number {n}", fields=fields, meta={"row": n})


@pytest.fixture
def results() -> ResultsFile:
    items = [
        _item(2, "bug", "jev", "Camera"),
        _item(3, "bug", "llm", "Camera"),
        _item(4, "feature_request", "jev", "Login", low=True),
        _item(5, "praise", "jev", "Login"),
    ]
    return ResultsFile(
        run_id="r_test0001",
        job="Sort reviews into bugs and requests",
        threshold=0.7,
        started_at="2026-10-11T08:00:00Z",
        finished_at="2026-10-11T08:00:12Z",
        plan={
            "nodes": [
                {"id": "n1", "tool": "classify", "params": {}},
                {"id": "n2", "tool": "group", "params": {}},
            ]
        },
        items=items,
        dropped=[DroppedItem(id="row-6", text="buy cheap pills", reason="spam")],
        aggregates={
            "type_counts": {"bug": 2, "feature_request": 1, "praise": 1},
            "theme_groups": [
                {"theme": "Camera", "description": "camera", "count": 2, "example_ids": ["row-2"]},
                {"theme": "Login", "description": "login", "count": 2, "example_ids": ["row-4"]},
            ],
            "theme_counts": {"Camera": 2, "Login": 2},
            "rank_order": ["row-3", "row-2"],
            "escalation_curve": [[0.0, 0], [0.5, 1], [1.0, 8]],
        },
        node_reports=[
            NodeReport(
                node="n1",
                tool="classify",
                items_in=5,
                items_out=4,
                jev_decisions=3,
                llm_decisions=1,
            ),
            NodeReport(node="n2", tool="group", items_in=4, items_out=4, jev_decisions=4),
        ],
        stats=RunStats(
            decisions_total=9,
            by={"jev": 7, "llm": 1, "code": 1},
            jev_cost_usd=0.0012,
            llm_cost_usd=0.0034,
            wall_time_s=12.4,
            jev_calls=4,
            llm_calls=1,
        ),
        decisions_file="runs/r_test0001/decisions.jsonl",
    )


async def _summarize(
    results: ResultsFile, text: str | None, fail: bool = False
) -> tuple[str, object]:
    responses = {} if text is None else {"summary": {"summary": text}}
    ctx = await make_context(fake=True, llm=FakeLLM(json_responses=responses, fail=fail))
    try:
        return await summarize(results, ctx), ctx
    finally:
        await ctx.aclose()


async def test_summary_with_input_numbers_is_kept(results):
    text = "Foreman sorted 4 reviews; 2 were bugs. Jev made 7 decisions and the LLM 1, for $0.0046."
    out, ctx = await _summarize(results, text)
    assert out == text
    assert invented_numbers(out, summary_input(results)) == []
    rows = [d for d in ctx.ledger.decisions if d.tool == "format"]
    assert rows and rows[0].kind == "proposal" and rows[0].decided_by == "llm"


async def test_summary_with_invented_number_falls_back(results):
    out, _ = await _summarize(results, "Foreman sorted 4,812 reviews and saved 37 hours.")
    assert out == template_summary(results)


async def test_summary_llm_failure_falls_back(results):
    out, _ = await _summarize(results, None, fail=True)
    assert out == template_summary(results)
    assert "Jev made 7 of the 8 model decisions" in out


def test_rounded_and_percent_numbers(results):
    src = summary_input(results)
    assert invented_numbers("Jev made 87.5% of model decisions in 12.4 s.", src) == []
    assert invented_numbers("That is 91% cheaper.", src) == ["91%"]


def test_charts_by_shape(results):
    charts = find_charts(results)
    keys = [c.key for c in charts]
    assert keys[0] == "decided_by" and charts[0].split
    assert "type_counts" in keys
    # theme_groups and theme_counts have the same bars: one chart only.
    assert ("theme_groups" in keys) != ("theme_counts" in keys)
    assert "escalation_curve" not in keys and "rank_order" not in keys


def test_xlsx_sheets_and_colours(results):
    import openpyxl

    wb = openpyxl.load_workbook(BytesIO(export(results, "xlsx", summary="Short summary.")))
    assert wb.sheetnames[:2] == ["Summary", "Results"]
    assert {"Type counts", "Theme groups", "Stats", "Dropped"} <= set(wb.sheetnames)
    ws = wb["Results"]
    header = [c.value for c in ws[2]]
    assert header[:2] == ["id", "text"]
    assert ws["A3"].value == "row-3"  # ranked first
    type_col = header.index("type") + 1
    llm_cell = ws.cell(row=3, column=type_col)  # row-3 was decided by the LLM
    jev_cell = ws.cell(row=4, column=type_col)  # row-2 by Jev
    assert llm_cell.fill.fgColor.rgb.endswith(FILL["llm"])
    assert jev_cell.fill.fgColor.rgb.endswith(FILL["jev"])
    low_row = next(r for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=1).value == "row-4")
    assert ws.cell(row=low_row, column=type_col).font.color.rgb.endswith("D70015")
    hidden = header.index("type__by") + 1
    assert ws.column_dimensions[ws.cell(row=2, column=hidden).column_letter].hidden
    assert ws["B1"].value == "Jev"  # legend row


def test_pdf_and_docx(results):
    pdf = export(results, "pdf", summary="Short summary.")
    assert pdf.startswith(b"%PDF")
    import docx

    d = docx.Document(BytesIO(export(results, "docx")))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Foreman run report" in text and "Type counts" in text


def test_unknown_format(results):
    with pytest.raises(ValueError):
        export(results, "csv")  # type: ignore[arg-type]
