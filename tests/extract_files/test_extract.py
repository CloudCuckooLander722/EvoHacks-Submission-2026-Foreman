"""extract(): every format from fixtures built in code, plus the edge cases."""

import json
from datetime import date, datetime
from pathlib import Path

import pytest

from foreman.contracts import ItemsFile
from foreman.core.context import make_context
from foreman.extract_files import ExtractError, detect_kind, extract
from foreman.extract_files.__main__ import main as cli
from foreman.extract_files.check import check_readability
from foreman.extract_files.text import MAX_SECTION_CHARS

REVIEWS = [
    ("App crashes every time I open the camera on my phone.", 1, "2026-09-30"),
    ("Please add a dark mode, my eyes hurt at night when reading.", 3, "2026-10-01"),
    ("Love it, the best reading app I have used in years.", 5, "2026-10-02"),
]


def _valid(items: ItemsFile) -> ItemsFile:
    return ItemsFile.model_validate(json.loads(items.model_dump_json()))


def test_csv_rows_typed_fields(tmp_path: Path):
    p = tmp_path / "reviews.csv"
    lines = ["review,rating,date,zip,empty"] + [f'"{t}",{r},{d},02134,' for t, r, d in REVIEWS]
    p.write_text("\n".join(lines), encoding="utf-8")
    out = _valid(extract(p))
    assert [i.id for i in out.items] == ["row-2", "row-3", "row-4"]
    first = out.items[0]
    assert first.text == REVIEWS[0][0]
    assert first.fields == {
        "review": REVIEWS[0][0],
        "rating": 1,
        "date": "2026-09-30",
        "zip": "02134",
    }
    assert first.meta == {"row": 2}
    assert out.source["text_column"] == "review"
    assert out.source["kind"] == "csv"
    assert out.source["rows"] == 3


def test_csv_without_text_column_joins_row(tmp_path: Path):
    p = tmp_path / "nums.csv"
    p.write_text("a,b\n1,2.5\n3,4\n", encoding="utf-8")
    out = extract(p)
    assert out.items[0].text == "a: 1; b: 2.5"
    assert out.items[0].fields == {"a": 1, "b": 2.5}
    assert out.source["text_column"] is None


def test_cp1252_csv(tmp_path: Path):
    p = tmp_path / "latin.csv"
    p.write_bytes('name,comment\ncafé,"très bien, merci beaucoup pour l\'app"\n'.encode("cp1252"))
    out = extract(p)
    assert out.items[0].fields["name"] == "café"


def test_xlsx_first_sheet_dates(tmp_path: Path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feedback"
    ws.append(["review", "rating", "when", "score"])
    for t, r, d in REVIEWS:
        ws.append([t, r, datetime.fromisoformat(d), 1.5])
    ws.append([None, None, None, None])  # trailing formatted-but-empty row
    wb.create_sheet("Other").append(["x"])
    p = tmp_path / "reviews.xlsx"
    wb.save(p)

    assert detect_kind(p) == "xlsx"
    out = _valid(extract(p))
    assert [i.id for i in out.items] == ["row-2", "row-3", "row-4"]
    assert out.items[1].fields == {
        "review": REVIEWS[1][0],
        "rating": 3,
        "when": "2026-10-01",
        "score": 1.5,
    }
    assert out.items[1].meta == {"row": 3, "sheet": "Feedback"}
    assert out.source["sheet"] == "Feedback"
    assert out.source["text_column"] == "review"
    json.dumps(out.model_dump())  # no NaN / Timestamp / datetime left


def test_json_list_uses_ids_and_text_key(tmp_path: Path):
    p = tmp_path / "tickets.json"
    p.write_text(
        json.dumps([{"id": "T1", "body": "Login fails", "n": 2}, {"id": "T2", "body": "Slow"}])
    )
    out = extract(p)
    assert [i.id for i in out.items] == ["T1", "T2"]
    assert out.items[0].text == "Login fails"
    assert out.items[0].fields == {"id": "T1", "body": "Login fails", "n": 2}


def test_json_duplicate_ids_fall_back(tmp_path: Path):
    p = tmp_path / "dups.json"
    p.write_text(json.dumps([{"id": 1, "note": "aaa"}, {"id": 1, "note": "bbbb"}]))
    out = extract(p)
    assert [i.id for i in out.items] == ["item-1", "item-2"]
    assert out.items[1].text == "bbbb"


def test_items_json_passes_through():
    out = extract("data/samples/reviews_small.json")
    assert out.items[0].id == "row-2"


def test_bad_items_json_raises(tmp_path: Path):
    p = tmp_path / "bad.json"
    p.write_text(json.dumps({"items": [{"id": "a", "text": 5}]}))
    with pytest.raises(ExtractError, match="not a valid items.json"):
        extract(p)


def test_txt_list_one_item_per_line(tmp_path: Path):
    p = tmp_path / "notes.txt"
    p.write_text("\n".join(f"note number {n}" for n in range(1, 7)) + "\n\n", encoding="utf-8")
    out = extract(p)
    assert [i.id for i in out.items][:2] == ["line-1", "line-2"]
    assert len(out.items) == 6


def test_txt_prose_packed_into_sections(tmp_path: Path):
    para = ("This is a long sentence about the product and its many problems. " * 60).strip()
    p = tmp_path / "essay.txt"
    p.write_text("\n\n".join([para] * 12), encoding="utf-8")
    out = extract(p)
    assert len(out.items) >= 2
    assert all(len(i.text) <= MAX_SECTION_CHARS for i in out.items)
    assert out.items[0].id == "essay-s1"


def test_pdf_pages_into_sections(tmp_path: Path):
    import pymupdf

    doc = pymupdf.open()
    for n in range(1, 4):
        doc.new_page().insert_text((72, 72), f"Page {n} says hello.")
    p = tmp_path / "report.pdf"
    doc.save(p)
    out = extract(p)
    assert len(out.items) == 1
    assert out.items[0].id == "report-p1-s1"
    assert out.items[0].meta == {"page_start": 1, "page_end": 3}
    assert "Page 3 says hello." in out.items[0].text


def test_docx_sections_by_heading(tmp_path: Path):
    import docx

    d = docx.Document()
    d.add_paragraph("Intro line before any heading.")
    d.add_heading("Bugs", level=1)
    d.add_paragraph("The camera crashes.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "id", "issue"
    t.cell(1, 0).text, t.cell(1, 1).text = "7", "login"
    d.add_heading("Requests", level=1)
    d.add_paragraph("Dark mode please.")
    p = tmp_path / "spec.docx"
    d.save(p)
    out = extract(p)
    assert [i.id for i in out.items] == ["spec-s1", "spec-s2", "spec-s3"]
    assert out.items[1].meta == {"section": 2, "heading": "Bugs"}
    assert "id: 7; issue: login" in out.items[1].text
    assert all(len(i.text) <= MAX_SECTION_CHARS for i in out.items)


def test_kind_from_bytes_not_extension(tmp_path: Path):
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "hidden pdf")
    p = tmp_path / "renamed.txt"
    doc.save(p)
    assert detect_kind(p) == "pdf"


@pytest.mark.parametrize(
    ("setup", "message"),
    [
        (lambda d: None, "No file path"),
        (lambda d: d / "missing.csv", "File not found"),
        (lambda d: d, "Not a regular file"),
        (lambda d: _write(d / "empty.csv", b""), "is empty"),
        (lambda d: _write(d / "blank.txt", b"   \n\n  "), "only whitespace"),
        (lambda d: _write(d / "bin.dat", b"\x00\x01\x02garbage"), "binary"),
        (lambda d: _write(d / "bad.json", b"{not json"), "not valid JSON"),
        (lambda d: _write(d / "x.zip", _zip_bytes()), "not an xlsx or docx"),
        (lambda d: _write(d / "noitems.csv", b"a,b\n"), "no items"),
    ],
)
def test_edge_cases_raise(tmp_path: Path, setup, message):
    with pytest.raises(ExtractError, match=message):
        extract(setup(tmp_path))


def _write(p: Path, data: bytes) -> Path:
    p.write_bytes(data)
    return p


def _zip_bytes() -> bytes:
    import io
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "hi")
    return buf.getvalue()


async def test_readability_check_offline():
    items = extract("data/samples/reviews_small.json")
    ctx = await make_context(fake=True)
    try:
        warnings = await check_readability(items, ctx, sample=5)
        assert isinstance(warnings, list)
        assert ctx.ledger.decisions and ctx.ledger.decisions[0].tool == "extract_files"
    finally:
        await ctx.aclose()


def test_cli_writes_items(tmp_path: Path, capsys):
    src = tmp_path / "r.csv"
    src.write_text("review,rating\n" + "\n".join(f'"{t}",{r}' for t, r, _ in REVIEWS))
    out = tmp_path / "items.json"
    assert cli([str(src), "-o", str(out)]) == 0
    assert len(ItemsFile.model_validate_json(out.read_text()).items) == 3
    assert cli([str(tmp_path / "nope.csv"), "-o", str(out)]) == 1


def test_dates_survive(tmp_path: Path):
    from foreman.extract_files.tables import clean_value

    assert clean_value(date(2026, 1, 2)) == "2026-01-02"
    assert clean_value(float("nan")) is None
